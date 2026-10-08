"""
Unit tests for Excel and CSV export functionality.
"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import openpyxl
import pandas as pd
from src.pipeline import StatementPipeline
from src.exporters import ExcelExporter, CSVExporter

SAMPLE_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "sample_data")


class TestExport(unittest.TestCase):

    def setUp(self):
        self.pipeline = StatementPipeline()
        pdf_path = os.path.join(SAMPLE_DIR, "hdfc_sample_statement.pdf")
        self.result = self.pipeline.process(pdf_path)

    def test_excel_export(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            ExcelExporter.export(self.result, tmp_path)
            self.assertTrue(os.path.exists(tmp_path))

            wb = openpyxl.load_workbook(tmp_path)
            expected_sheets = ["Account Overview", "Classified Transactions", "Category Breakdown", "Monthly Trends"]
            for s in expected_sheets:
                self.assertIn(s, wb.sheetnames)

            # Check rows in Transactions sheet
            ws = wb["Classified Transactions"]
            self.assertGreater(ws.max_row, 10)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_csv_export(self):
        with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            CSVExporter.export_result(self.result, tmp_path)
            self.assertTrue(os.path.exists(tmp_path))

            df = pd.read_csv(tmp_path)
            self.assertEqual(len(df), len(self.result.transactions))
            expected_cols = ["Date", "Description", "Debit", "Credit", "Balance", "Category"]
            for col in expected_cols:
                self.assertIn(col, df.columns)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


if __name__ == "__main__":
    unittest.main()
