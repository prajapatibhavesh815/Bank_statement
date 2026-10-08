"""
Professional Multi-Sheet Excel Exporter (.xlsx).
Produces styled workbooks with Account Summary, Classified Transactions,
Category Breakdown, and Monthly Trends.
"""
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ..models.schemas import StatementProcessingResult, Transaction, AccountDetails
from ..analytics.financial_summary import FinancialSummary


class ExcelExporter:
    """
    Exports processing results into a beautifully formatted Excel document.
    """

    HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # Navy Blue
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F497D")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True)
    REGULAR_FONT = Font(name="Calibri", size=11)
    
    CREDIT_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Soft Green
    DEBIT_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")   # Soft Red

    THIN_BORDER = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    @classmethod
    def export(cls, result: StatementProcessingResult, file_path: str):
        """
        Creates the complete Excel workbook and saves to file_path.
        """
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # 1. Sheet: Account Overview
        cls._create_overview_sheet(wb, result)

        # 2. Sheet: Classified Transactions
        cls._create_transactions_sheet(wb, result.transactions)

        # 3. Sheet: Category Breakdown
        analytics = FinancialSummary.analyze(result.transactions, result.account_details)
        cls._create_category_sheet(wb, analytics["category_summary"])

        # 4. Sheet: Monthly Trends
        cls._create_monthly_sheet(wb, analytics["monthly_trend"])

        wb.save(file_path)

    @classmethod
    def _create_overview_sheet(cls, wb: openpyxl.Workbook, result: StatementProcessingResult):
        ws = wb.create_sheet(title="Account Overview")
        ws.views.sheetView[0].showGridLines = True

        acc = result.account_details
        meta = result.metadata

        # Title
        ws["A1"] = "BANK STATEMENT AUDIT & CLASSIFICATION SUMMARY"
        ws["A1"].font = cls.TITLE_FONT

        # Account Details Section
        info_rows = [
            ("Bank Name", acc.bank_name),
            ("Account Holder", acc.account_holder),
            ("Account Number", acc.account_number),
            ("IFSC Code", acc.ifsc_code or "N/A"),
            ("Branch", acc.branch or "N/A"),
            ("Statement Period", f"{acc.statement_period_start or 'N/A'} to {acc.statement_period_end or 'N/A'}"),
            ("PDF Detection Type", meta.file_type),
            ("Total Transactions", meta.total_transactions),
            ("Total Debits (Withdrawals)", f"₹ {meta.total_debits:,.2f}"),
            ("Total Credits (Deposits)", f"₹ {meta.total_credits:,.2f}"),
            ("Net Cash Flow", f"₹ {meta.net_cash_flow:,.2f}"),
            ("Opening Balance", f"₹ {acc.opening_balance:,.2f}" if acc.opening_balance is not None else "N/A"),
            ("Closing Balance", f"₹ {acc.closing_balance:,.2f}" if acc.closing_balance is not None else "N/A"),
            ("Balance Audit Status", "100% Reconciled" if meta.balance_mismatches_count == 0 else f"{meta.balance_mismatches_count} Discrepancies Flagged"),
        ]

        ws.append([])  # Row 2 empty
        start_row = 3
        for label, val in info_rows:
            cell_lbl = ws.cell(row=start_row, column=1, value=label)
            cell_val = ws.cell(row=start_row, column=2, value=val)
            cell_lbl.font = cls.BOLD_FONT
            cell_lbl.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            cell_val.font = cls.REGULAR_FONT
            cell_lbl.border = cls.THIN_BORDER
            cell_val.border = cls.THIN_BORDER
            start_row += 1

        cls._auto_fit_columns(ws)

    @classmethod
    def _create_transactions_sheet(cls, wb: openpyxl.Workbook, transactions: List[Transaction]):
        ws = wb.create_sheet(title="Classified Transactions")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Date", "Description", "Merchant / Entity", "Channel",
            "Debit (Dr)", "Credit (Cr)", "Balance",
            "Assigned Category", "Sub-Category", "Confidence",
            "Classification Method", "Audit Verified"
        ]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 26

        for row_idx, t in enumerate(transactions, start=2):
            ws.cell(row=row_idx, column=1, value=t.date).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=t.description)
            ws.cell(row=row_idx, column=3, value=t.merchant or "-")
            ws.cell(row=row_idx, column=4, value=t.channel or "-").alignment = Alignment(horizontal="center")

            # Debit
            dr_cell = ws.cell(row=row_idx, column=5, value=t.debit if t.debit > 0 else "")
            dr_cell.number_format = "#,##0.00"
            dr_cell.alignment = Alignment(horizontal="right")
            if t.debit > 0:
                dr_cell.fill = cls.DEBIT_FILL

            # Credit
            cr_cell = ws.cell(row=row_idx, column=6, value=t.credit if t.credit > 0 else "")
            cr_cell.number_format = "#,##0.00"
            cr_cell.alignment = Alignment(horizontal="right")
            if t.credit > 0:
                cr_cell.fill = cls.CREDIT_FILL

            # Balance
            bal_cell = ws.cell(row=row_idx, column=7, value=t.balance)
            bal_cell.number_format = "#,##0.00"
            bal_cell.alignment = Alignment(horizontal="right")

            # Category & Metadata
            ws.cell(row=row_idx, column=8, value=t.category).font = cls.BOLD_FONT
            ws.cell(row=row_idx, column=9, value=t.sub_category or "-")
            conf_cell = ws.cell(row=row_idx, column=10, value=f"{int(t.confidence * 100)}%")
            conf_cell.alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=11, value=t.classification_method).alignment = Alignment(horizontal="center")
            
            # Audit
            audit_cell = ws.cell(row=row_idx, column=12, value="Verified" if t.balance_verified else "Mismatch")
            audit_cell.alignment = Alignment(horizontal="center")

            # Borders
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).border = cls.THIN_BORDER

        cls._auto_fit_columns(ws)

    @classmethod
    def _create_category_sheet(cls, wb: openpyxl.Workbook, categories: List[Dict[str, Any]]):
        ws = wb.create_sheet(title="Category Breakdown")
        ws.views.sheetView[0].showGridLines = True

        headers = ["Category", "Total Spent (Debit)", "Txn Count", "% of Expenses"]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        ws.row_dimensions[1].height = 24

        for row_idx, item in enumerate(categories, start=2):
            ws.cell(row=row_idx, column=1, value=item["category"]).font = cls.BOLD_FONT
            amt_cell = ws.cell(row=row_idx, column=2, value=item["total_amount"])
            amt_cell.number_format = "#,##0.00"
            amt_cell.alignment = Alignment(horizontal="right")

            cnt_cell = ws.cell(row=row_idx, column=3, value=item["count"])
            cnt_cell.alignment = Alignment(horizontal="center")

            pct_cell = ws.cell(row=row_idx, column=4, value=f"{item['percentage']}%")
            pct_cell.alignment = Alignment(horizontal="right")

            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).border = cls.THIN_BORDER

        cls._auto_fit_columns(ws)

    @classmethod
    def _create_monthly_sheet(cls, wb: openpyxl.Workbook, trends: List[Dict[str, Any]]):
        ws = wb.create_sheet(title="Monthly Trends")
        ws.views.sheetView[0].showGridLines = True

        headers = ["Month", "Total Debits (Expenses)", "Total Credits (Income)", "Net Savings / Flow"]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        ws.row_dimensions[1].height = 24

        for row_idx, item in enumerate(trends, start=2):
            ws.cell(row=row_idx, column=1, value=item["month"]).alignment = Alignment(horizontal="center")
            
            dr_cell = ws.cell(row=row_idx, column=2, value=item["total_debit"])
            dr_cell.number_format = "#,##0.00"
            dr_cell.alignment = Alignment(horizontal="right")

            cr_cell = ws.cell(row=row_idx, column=3, value=item["total_credit"])
            cr_cell.number_format = "#,##0.00"
            cr_cell.alignment = Alignment(horizontal="right")

            net_cell = ws.cell(row=row_idx, column=4, value=item["net"])
            net_cell.number_format = "#,##0.00"
            net_cell.alignment = Alignment(horizontal="right")
            if item["net"] >= 0:
                net_cell.fill = cls.CREDIT_FILL
            else:
                net_cell.fill = cls.DEBIT_FILL

            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).border = cls.THIN_BORDER

        cls._auto_fit_columns(ws)

    @classmethod
    def _auto_fit_columns(cls, ws: openpyxl.worksheet.worksheet.Worksheet):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)
